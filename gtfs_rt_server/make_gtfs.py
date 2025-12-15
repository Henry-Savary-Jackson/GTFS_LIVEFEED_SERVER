import pandas as pd
import openpyxl
import re
import zipfile
import sqlalchemy
import subprocess
import os
import io
import sys
import datetime
import json
import numpy as np

FILE_PATH = os.path.dirname(__file__)


def has_errors(result_path):
    with open(os.path.normpath(f"{result_path}/report.json"), "r") as file:
        report_json = json.load(file)
        notices = report_json["notices"]
        print(notices)
        for notice in notices:
            if notice["severity"] in ["ERROR", "WARNING"]:
                return True
        return False


def validate_gtfs(validator_path, zipfile_path, result_path, update_method=None):
    with subprocess.Popen(
        [
            "java",
            "-jar",
            validator_path,
            "--input",
            zipfile_path,
            "-o",
            result_path,
        ],
        stderr=subprocess.PIPE,
        stdout=subprocess.PIPE,
        text=True,
    ) as process:
        print("Waiting for jar file validation")
        if update_method:
            update_method(status="working", message="Waiting for jar file validation")
        stdout, stderr = process.communicate()
        if process.returncode != 0:
            raise Exception(f"error:\n{stderr}")
        else:
            print(stdout)
            if update_method:
                update_method(status="working", message=stdout or "")


def read_sheet_as_df(worksheet, **kwargs):
    try:
        start = kwargs["skiprows"] + 1 if "skiprows" in kwargs else 1
        nrows = kwargs["nrows"] if "" in kwargs else None
        start_cell = worksheet.cell(start, 1)
        current_cell = start_cell
        columns = []
        i = 1
        while current_cell.value not in ["", " ", None]:
            value = current_cell.value
            if type(value) == float:  # remove decimals
                value = int(value)
            columns.append(str(value))
            i += 1
            current_cell = worksheet.cell(current_cell.row, i)

        row = start + 1
        current_cell = worksheet.cell(row, 1)
        rows = []
        while current_cell.value not in ["", " ", None, "TRAIN NO.", "Service"] and (
            not nrows or row - start <= nrows - 1
        ):
            # add
            current_row = []
            for col in range(1, len(columns) + 1):
                current_cell = worksheet.cell(row, col)
                current_row.append(current_cell.value)
            rows.append(current_row)
            row += 1
            current_cell = worksheet.cell(row, 1)

        df = pd.DataFrame(rows, columns=columns)

        return df
    except ValueError as e:
        if "Excel file format cannot be determined" in str(e):
            print(str(e))
            raise Exception("Invalid Excel file")
        raise e


def getDistMatrixDataFrame(workbook):
    return read_sheet_as_df(workbook["Distance Matrix"])

def getTicketTypeDataframe(workbook):
    return read_sheet_as_df(workbook["Ticket prices"])

def getKmZoneDataframe(workbook):
    return read_sheet_as_df(workbook["Km Zones"])

def getRoutesDataFrame(workbook):
    df = read_sheet_as_df(workbook["Routes"])
    df["route_type"] = df["route_type"].astype(np.int64)
    return df


def getFareRulesDataFrame(workbook):
    return read_sheet_as_df(workbook["FareRules"])


def getFareAttributesDataFrame(workbook):
    df = read_sheet_as_df(workbook["FareAttributes"])
    if "transfers" in df:
        print(df["transfers"])
        df["transfers"] = df["transfers"].fillna(-1).astype(np.int64).astype(str)
        df["transfers"] = df["transfers"].replace("-1", "")
    df["payment_method"] = df["payment_method"].astype(np.int64)
    return df


def getStopsDataFrame(workbook):
    stops_df = read_sheet_as_df(workbook["Stops"])
    stops_df["zone_id"] = stops_df["stop_name"]
    return stops_df


def getServicesDataFrame(workbook):
    df = read_sheet_as_df(workbook["Services"])
    df["start_date"] = df["start_date"].astype(np.int64)
    df["end_date"] = df["end_date"].astype(np.int64)
    df["monday"] = df["monday"].astype(np.int64)
    df["tuesday"] = df["tuesday"].astype(np.int64)
    df["wednesday"] = df["wednesday"].astype(np.int64)
    df["thursday"] = df["thursday"].astype(np.int64)
    df["friday"] = df["friday"].astype(np.int64)
    df["saturday"] = df["saturday"].astype(np.int64)
    df["sunday"] = df["sunday"].astype(np.int64)
    return df


def getShapesDataFrame(workbook):
    df = read_sheet_as_df(workbook["Shapes"])
    df["shape_pt_sequence"] = df["shape_pt_sequence"].astype(np.int64)
    return df


def getAgencyDataFrame(workbook):
    return read_sheet_as_df(workbook["AgencyInfo"])


def getFeedInfoDataFrame(workbook):
    df = read_sheet_as_df(workbook["FeedInfo"])
    df["feed_start_date"] = df["feed_start_date"].astype(np.int64)
    df["feed_end_date"] = df["feed_end_date"].astype(np.int64)
    return df


def getCalendarDaysDataFrame(workbook):
    df = read_sheet_as_df(workbook["CalendarDays"])
    return df


def getStops(stops_df):
    return set(stops_df["stop_id"])


def getShapes(shapes_df):
    return set(shapes_df["shape_id"])


def getRoutes(routes_df):
    return set(routes_df["route_id"])


def getServices(services_df):
    return set(services_df["service_id"])


def get_metadata(worksheet, sheet_name, services, shapes, start):
    df = read_sheet_as_df(worksheet, skiprows=start - 1, nrows=2)

    if "Service" not in df.columns:
        raise ValueError(f"No Service column provided in sheet {sheet_name}")
    if df["Service"][0] not in services:
        raise ValueError(
            f"No such service ( {df["Service"][0] or ""} ) in sheet {sheet_name}  "
        )

    if "Shape" not in df.columns:
        raise ValueError(
            f"Your shape column is missing or in the incorrect place at the top of the schedule."
        )

    if pd.notna(df["Shape"][0]) and df["Shape"][0] not in shapes:
        raise ValueError(f"No such shape ( {df["Shape"][0]} ) in sheet {sheet_name}  ")

    return df["Service"][0], df["Shape"][0]


def read_schedule_as_df(worksheet, start, length):
    return read_sheet_as_df(worksheet, nrows=length, skiprows=start - 1)


def get_timetable_info(
    worksheet,
    sheet_name,
    start,
    length,
    sheet_title_directory,
    stoptime_df: pd.DataFrame,
    stop_df,
    trip_df,
    sub_routes,
    route_id,
    services,
    shapes,
    stops,
):
    service_id, shape_id = get_metadata(
        worksheet, sheet_title_directory, services, shapes, start
    )

    df_schedule = read_schedule_as_df(worksheet, start + 2, length)
    if "TRAIN NO." not in df_schedule.columns:
        raise ValueError(f'TRAIN NO. not in the sheet "{sheet_title_directory}".')
    df_schedule = df_schedule.set_index("TRAIN NO.").dropna(axis=1, how="all")
    # keep a list of trains for each sub_route
    for index_col, train_number in enumerate(df_schedule.columns):

        ## if train_number
        train_number = str(train_number)
        new_route_id = f"{route_id}-{train_number[:2] if len(train_number) >= 4 else train_number[:1]}s"
        sub_routes.add(new_route_id)

        dot_index = train_number.rfind(".")
        if dot_index != -1:
            train_number = train_number[:dot_index]

        stop_times_trip_df = pd.DataFrame(columns=stoptime_df.columns)
        trip_id = f"{service_id}-{train_number}"
        # must make list so that string comparison works
        if trip_id in list(trip_df["trip_id"]):
            raise ValueError(
                f'Duplicate trip "{trip_id}" in sheet {sheet_title_directory}.'
            )

        stop = ""
        skip = 0
        for i in range(len(df_schedule.index)):
            time = df_schedule.iloc[i, index_col]
            # if None, Nan or "..", skip adding stop time
            if not time or pd.isna(time) or time == "..":
                skip += 1
                continue
            # convert to str if it is time object
            if type(time) == datetime.time:
                time = time.strftime("%H:%M:%S")

            if time.rfind(":") == time.find(":"):
                time += ":00"
            stop = df_schedule.index[i]
            if stop not in stops:
                raise ValueError(
                    f'Stop "{stop}" in sheet {sheet_title_directory} doesn\'t exist.'
                )
            stop_times_trip_df = stop_times_trip_df._append(
                {
                    "trip_id": trip_id,
                    "arrival_time": time,
                    "departure_time": time,
                    "stop_id": stop,
                    "stop_sequence": i - skip,
                    "timepoint": 1,
                },
                ignore_index=True,
            )
        stoptime_df = pd.concat([stoptime_df, stop_times_trip_df])
        trip_df = trip_df._append(
            {
                "route_id": new_route_id,
                "service_id": service_id,
                "trip_id": trip_id,
                "trip_headsign": get_stop_name(stop_df, stop),
                "shape_id": shape_id or "",
            },
            ignore_index=True,
        )
    return stoptime_df, trip_df


def add_schedule(
    worksheet,
    sheet_name,
    sheet_title_directory,
    stoptime_df: pd.DataFrame,
    stop_df,
    trip_df,
    sub_routes,
    route_id,
    services,
    shapes,
    stops,
):
    try:
        min_row = worksheet.min_row
        max_row = worksheet.max_row
        current_row = min_row
        while current_row <= max_row:
            current_cell = worksheet.cell(current_row, 1)
            if current_cell.value == "Service":
                current_length = 0
                row_for_schedule = (
                    current_row + 2
                )  # increment current row until you reach the  end of the schedule
                current_cell = worksheet.cell(row_for_schedule, 1)
                while (
                    current_cell.value not in ["", " ", None, "Service"]
                    and row_for_schedule <= max_row
                ):
                    current_length += 1
                    row_for_schedule += 1
                    current_cell = worksheet.cell(row_for_schedule, 1)

                stoptime_df, trip_df = get_timetable_info(
                    worksheet,
                    sheet_name,
                    current_row,
                    current_length,
                    sheet_title_directory,
                    stoptime_df,
                    stop_df,
                    trip_df,
                    sub_routes,
                    route_id,
                    services,
                    shapes,
                    stops,
                )
                current_row = row_for_schedule
                continue

            current_row += 1
        print(stoptime_df.shape)
        return stoptime_df, trip_df
    except ValueError as e:

        raise e


def handle_km_zones(km_zones_df):
    return km_zones_df["distance"].tolist()

def handle_ticket_types(ticket_types_df, km_zones):
    ticket_types_dict = {}
    for i, ticket_types_row in ticket_types_df.iterrows():
        ticket_type = ticket_types_row["ticket_type"]
        ticket_price = ticket_types_row["price"]
        ticket_kmzone = int(ticket_types_row["km_zone"])
        if ticket_type in ticket_types_dict:
            ticket_types_dict[ticket_type][ticket_kmzone-1] = ticket_price
        else:
            new_list = [0 for km_zone in km_zones]
            new_list[ticket_kmzone-1] = ticket_price 
            ticket_types_dict[ticket_type] = new_list

    return ticket_types_dict





def handle_dist_matrix( dist_matrix, distance_prices, ticket_types):
    df_fare_rules = []
    df_fare_attributes = []
    for i in range(1, len(dist_matrix.columns)):
        # get all distance pairs
        origin_name = dist_matrix.columns[i].strip()
        for j in range(i, len(dist_matrix)):
            dest_name = dist_matrix["Stops"].iloc[j].strip()
            distance = dist_matrix.iloc[j,i]
            price_index = np.argmax(distance_prices > distance)
            # for each set of tickets get the relevant price

            for ticket_type in ticket_types:
                price = ticket_types[ticket_type][price_index]
                origin_name =origin_name.strip()
                fare_id_1 = f"{ticket_type}-kmzone{price_index+1}-{origin_name}-{dest_name}"
                fare_id_2 = f"{ticket_type}-kmzone{price_index+1}-{dest_name}-{origin_name}"

                df_fare_attributes.append([fare_id_1, price, "R", "0"])
                df_fare_attributes.append([fare_id_2, price, "R", "0"])
                df_fare_rules.append([fare_id_1, origin_name, dest_name])
                df_fare_rules.append([fare_id_2, dest_name, origin_name])

            # add row to fare_rules
            # add row to fare_attributes
            # add to discount

    df_fare_rules = pd.DataFrame(df_fare_rules, columns=["fare_id" ,	"origin_id" ,	"destination_id"])
    df_fare_attributes = pd.DataFrame(df_fare_attributes, columns=["fare_id" ,	"price" ,	 	"currency_type", 	"payment_method"])

    return  df_fare_rules, df_fare_attributes


def generate_gtfs_zip(
    excel_file, export_location, validator_path, result_path, update_method=None
):
    trip_df = pd.DataFrame(
        columns=["route_id", "service_id", "trip_id", "trip_headsign", "shape_id"]
    )
    stop_time_df = pd.DataFrame(
        columns=[
            "trip_id",
            "arrival_time",
            "departure_time",
            "stop_id",
            "stop_sequence",
            "timepoint",
        ]
    )
    if update_method:
        update_method(status="working", message="Reading Spreadsheets")
    print("getting dataframes")

    workbook = openpyxl.load_workbook(excel_file, data_only=True)
    routes_df = getRoutesDataFrame(workbook)
    services_df = getServicesDataFrame(workbook)
    shapes_df = getShapesDataFrame(workbook)
    agency_df = getAgencyDataFrame(workbook)
    calendar_days_df = getCalendarDaysDataFrame(workbook)
    feed_info_df = getFeedInfoDataFrame(workbook)
    stops_df = getStopsDataFrame(workbook)
    stops = getStops(stops_df)
    routes = getRoutes(routes_df)
    services = getServices(services_df)
    shapes = getShapes(shapes_df)

    km_zone_df = getKmZoneDataframe(workbook)
    ticket_types_df = getTicketTypeDataframe(workbook)
    distance_matrix_df  = getDistMatrixDataFrame(workbook)
    km_zones = handle_km_zones(km_zone_df)
    ticket_types_dict = handle_ticket_types(ticket_types_df,km_zones)
    fare_rules_df , fare_attributes_df = handle_dist_matrix(distance_matrix_df, km_zones, ticket_types_dict)

    new_routes = pd.DataFrame(columns=routes_df.columns)

    ## in binary mode right?
    print("getting directory")
    error = False
    directory = workbook["Directory"]
    for route in directory.iter_cols(min_col=2):
        sub_routes = set()

        route_name = route[0].value.strip()
        # get route id
        route_id = routes_df[routes_df["route_long_name"] == route_name].iloc[0][
            "route_id"
        ]
        for sheet_cell in route[1:]:
            if not sheet_cell.hyperlink:
                continue
            sheet_title_directory = sheet_cell.value
            try:
                sheet_name = get_sheet_name_from_hyperlink(
                    sheet_cell.hyperlink.location
                )

                if sheet_name not in workbook:
                    raise Exception(
                        f"The link for sheet '{sheet_title_directory}' in the Directory is an invalid link or that sheet doesnt exist."
                    )
                stop_time_df, trip_df = add_schedule(
                    workbook[sheet_name],
                    sheet_name,
                    sheet_title_directory,
                    stop_time_df,
                    stops_df,
                    trip_df,
                    sub_routes,
                    route_id,
                    services,
                    shapes,
                    stops,
                )
                print(sheet_name)
                if update_method:
                    update_method(message=f"Added {sheet_title_directory}")
            except Exception as e:
                print(e)
                print("error", sheet_name, e)
                if update_method:
                    update_method(
                        status="error-cont",
                        message=f'\nError in "{sheet_title_directory}":\n{e}\n',
                    )
                error = True

        # add sub_routes
        for sub_route in sub_routes:
            new_row = routes_df[routes_df["route_long_name"] == route_name].iloc[0]
            if len(sub_routes) > 1:
                new_row["route_long_name"] = (
                    f"{route_name}-{sub_route[sub_route.index("-")+1:]}"
                )
            new_row["route_id"] = sub_route
            new_routes = pd.concat([new_routes, pd.DataFrame(new_row).T], axis=0)


    df_dict = {
        "stop_times.txt": stop_time_df,
        "stops.txt": stops_df,
        "trips.txt": trip_df,
        "feed_info.txt": feed_info_df,
        "agency.txt": agency_df,
        "routes.txt": new_routes,
        "calendar_dates.txt": calendar_days_df,
        "shapes.txt": shapes_df,
        "calendar.txt": services_df,
        "fare_attributes.txt": fare_attributes_df,
        "fare_rules.txt": fare_rules_df,
    }

    if error:
        # there was an error in reading the gtfs zip
        raise Exception("There was at least 1 error in reading the excel file.")

    if update_method:
        update_method(message=f"Writing the zip file")

    try:
        with zipfile.ZipFile(export_location, "w") as gtfs_zip:
            for filename in df_dict:
                write_df_to_zipfile(gtfs_zip, filename, df_dict[filename])

    except Exception as e:
        if update_method:
            update_method(status="error", message=f"Error creating zip file: {e}")
        print(e)
        raise e

    try:
        print(validator_path)
        print(export_location)
        validate_gtfs(
            validator_path, export_location, result_path, update_method=update_method
        )
    except Exception as e:
        if update_method:
            update_method(status="error", message=f"Error validating zip file: {e}")
        print(e)
        raise e

    return df_dict


def add_gtfs_tables_to_db(engine: sqlalchemy.Engine, df_dict):
    for tablename in df_dict:
        df: pd.DataFrame = (
            df_dict[tablename]
            .dropna(axis=1, how="all")
            .astype("str")
            .replace("nan", pd.NA)
        )
        if df.empty:
            continue
        df.to_sql(
            tablename[: tablename.find(".")], engine, if_exists="replace", index=False
        )


def get_stop_name(stops_df, stop_id):
    return stops_df[stops_df["stop_id"] == stop_id].iloc[0]["stop_name"]


def get_sheet_name_from_hyperlink(link):
    if link.startswith("#"):
        link = link.split("#")[1].split("!")[0]
    link = link.split("!")[0]
    if link.startswith("'"):
        link = link[1 : len(link) - 1]
    return link


def write_df_to_zipfile(zip_file, filename, df):
    data_stream = io.StringIO()
    with data_stream:
        df.to_csv(data_stream, index=False)
        zip_file.writestr(zipfile.ZipInfo(filename), data_stream.getvalue())


if __name__ == "__main__":
    generate_gtfs_zip(
        open("/home/hsj/Downloads/new_sched.xlsx", "rb"),
        "./gtfs.zip",
        "./server_files/gtfs-validator-6.0.0-cli.jar",
        "server_files/shared_private/result",
    )
